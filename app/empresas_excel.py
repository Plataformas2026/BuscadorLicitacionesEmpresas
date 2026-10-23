"""
empresas_excel.py
-----------------
Carga de la BASE DE DATOS DE EMPRESAS (única fuente de verdad del sistema de recomendación de la Pestaña 2).

Origen: el Excel de empresas en GOOGLE DRIVE (GOOGLE_DRIVE_FILE_ID + GOOGLE_SERVICE_ACCOUNT_JSON, los mismos secretos que ya
usaba la sincronización anterior; ver `descargar_de_drive`). Si esos secretos no están configurados se usa la copia local
`data/empresas.xlsx` (ruta configurable con EMPRESAS_XLSX). Hoja «Hoja1», con la cabecera
en la fila 5 y una empresa por fila. Las columnas se localizan por el TEXTO de su cabecera (no por la letra), de modo que
añadir o mover columnas no rompe la carga; si falta alguna cabecera imprescindible se lanza un error que dice cuál. Las
demás hojas del libro (listas de sectores y un registro suelto) no se usan.

CAMPOS REALES DEL EXCEL (cabecera -> clave interna)
  EMPRESA (1ª columna con ese nombre)                 -> nombre_empresa      [campo 1 del ranking por embeddings]
  SECTOR                                              -> sector              [campo 2]
  palabras clave extraidas de LA DESCRIPCIÓN DE LA
  EMPRESA, PALABRAS CLAVE Y PROYECTOS TIPO de la bbdd -> palabras_clave      [campo 3]
  PRINCIPALES PROYECTOS TIPO                          -> proyectos_tipo      [campo 4]
  DESCRIPCIÓN DEL SECTOR                              -> descripcion_sector
  ESPECIALIDAD EXTRAIDA DE LAS REFERENCIAS            -> especialidad_referencias
  DESCRIPCIÓN BREVE DE LA ACTIVIDAD QUE DESARROLLA    -> descripcion_actividad
  PALABRAS CLAVE DE LAS REFERENCIAS / DE LA BBDD /
  palabras clave unificadas                           -> palabras_clave_referencias / _bbdd / _unificadas
  EXPERIENCIA PAÍSES                                  -> experiencia_paises
  ZONA GEOGRÁFICA y PAÍSES DE INTERÉS                 -> zona_geografica_interes
  ZONAS Y PAISES DESCARTADOS                          -> zonas_descartadas
  ÁMBITO GEOGRÁFICO DE OPERACIÓN                      -> ambito_geografico
  TAMAÑO / CERTIFICACIONES / FACTURACIÓN ANUAL /
  CLIENTES PRINCIPALES / importe mínimo / máximo      -> tamano / certificaciones / facturacion_anual /
                                                         clientes_principales / importes

El Excel NO trae un identificador de empresa: se genera `numero_interno` a partir del nombre (slug estable; si dos
nombres dieran el mismo se añade el número de fila). Los valores vacíos o «-» se cargan como None (nunca se inventa
nada). Este módulo no usa Streamlit ni Supabase.
"""
import io
import json
import os
import re
import unicodedata
from pathlib import Path

import openpyxl

HOJA = "Hoja1"
FILAS_BUSQUEDA_CABECERA = 20

# clave interna -> (texto normalizado con el que empieza la cabecera, requerida)
CABECERAS = {
    "nombre_empresa": ("empresa", True),
    "sector": ("sector", True),
    "descripcion_sector": ("descripcion del sector", False),
    "especialidad_referencias": ("especialidad extraida de las referencias", False),
    "descripcion_actividad": ("descripcion breve de la actividad", False),
    "palabras_clave_referencias": ("palabras clave de las referencias", False),
    "palabras_clave_bbdd": ("palabras clave de la bbdd", False),
    "palabras_clave_unificadas": ("palabras clave unificadas", False),
    "palabras_clave": ("palabras clave extraidas de la descripcion", True),
    "proyectos_tipo": ("principales proyectos tipo", True),
    "experiencia_paises": ("experiencia paises", False),
    "zona_geografica_interes": ("zona geografica y paises de interes", False),
    "zonas_descartadas": ("zonas y paises descartados", False),
    "tamano": ("tamano", False),
    "certificaciones": ("certificaciones", False),
    "facturacion_anual": ("facturacion anual", False),
    "clientes_principales": ("clientes principales", False),
    "ambito_geografico": ("ambito geografico de operacion", False),
    "importes": ("importe minimo", False),
}
VACIOS = {"", "-", "--", "—", "n/a", "na", "ninguno", "ninguna", "no consta", "s/d"}


def _norm(texto) -> str:
    texto = unicodedata.normalize("NFD", str(texto or "").casefold())
    return " ".join("".join(c for c in texto if unicodedata.category(c) != "Mn").split())


def limpiar(valor):
    """Texto limpio (conserva los saltos de línea del Excel) o None si está vacío."""
    if valor is None:
        return None
    lineas = [" ".join(linea.split()) for linea in str(valor).replace("\r", "\n").split("\n")]
    texto = "\n".join(l for l in lineas if l).strip()
    return None if _norm(texto) in VACIOS else texto


def a_linea(texto, separador: str = "; ") -> str:
    """Texto en una sola línea (los saltos de línea del Excel pasan a separadores)."""
    return separador.join(l.strip(" •·-•") for l in str(texto or "").split("\n") if l.strip(" •·-•"))


def trocear(texto, maximo: int = 60) -> list:
    """Lista de frases de un campo de palabras clave (separadas por coma, punto y coma, viñeta o salto de línea)."""
    vistas, resultado = set(), []
    for parte in re.split(r"[,;\n•·]+", str(texto or "")):
        parte = " ".join(parte.split()).strip(" .-–—:")
        clave = _norm(parte)
        if len(parte) >= 3 and clave not in vistas:
            vistas.add(clave)
            resultado.append(parte)
    return resultado[:maximo]


def _slug(nombre: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", _norm(nombre)).strip("-")[:60] or "empresa"


def ruta_por_defecto(ruta: str = None) -> Path:
    """Ruta del Excel: la indicada (absoluta o relativa a la raíz del proyecto) o `data/empresas.xlsx`."""
    raiz = Path(__file__).resolve().parent.parent
    ruta = ruta or "data/empresas.xlsx"
    candidata = Path(ruta) if os.path.isabs(ruta) else raiz / ruta
    if not candidata.exists() and ruta == "data/empresas.xlsx":
        otros = sorted((raiz / "data").glob("*.xlsx")) if (raiz / "data").is_dir() else []
        if len(otros) == 1:  # un único Excel en data/: se usa aunque tenga otro nombre
            return otros[0]
    return candidata


def _servicio_drive(credenciales):
    from google.oauth2 import service_account
    from googleapiclient.discovery import build

    info = credenciales if isinstance(credenciales, dict) else json.loads(credenciales)
    cred = service_account.Credentials.from_service_account_info(
        info, scopes=["https://www.googleapis.com/auth/drive.readonly"])
    return build("drive", "v3", credentials=cred, cache_discovery=False)


def fecha_modificacion_drive(file_id: str, credenciales) -> str:
    """modifiedTime del Excel en Drive (barato: solo metadatos). Sirve para saber si hay que volver a descargarlo."""
    return _servicio_drive(credenciales).files().get(fileId=file_id, fields="modifiedTime").execute()["modifiedTime"]


def descargar_de_drive(file_id: str, credenciales) -> io.BytesIO:
    """Descarga el Excel de Drive con la cuenta de servicio (acepta .xlsx subido o Google Sheets, que se exporta a .xlsx)."""
    servicio = _servicio_drive(credenciales)
    meta = servicio.files().get(fileId=file_id, fields="name, mimeType").execute()
    if meta.get("mimeType") == "application/vnd.google-apps.spreadsheet":
        contenido = servicio.files().export_media(
            fileId=file_id, mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet").execute()
    else:
        contenido = servicio.files().get_media(fileId=file_id).execute()
    return io.BytesIO(contenido)


def cargar_empresas(ruta=None, hoja: str = HOJA) -> list:
    """
    Lista de empresas (dict) del Excel. `ruta` puede ser una ruta, un Path o un archivo en memoria (BytesIO, p. ej. el
    descargado de Drive). Lanza FileNotFoundError / ValueError con un mensaje claro si algo falla.
    """
    en_memoria = hasattr(ruta, "read")
    if not en_memoria:
        ruta = ruta_por_defecto(ruta) if not isinstance(ruta, Path) else ruta
    if not en_memoria and not ruta.exists():
        raise FileNotFoundError(
            f"No se encuentra el Excel de empresas en «{ruta}». Copia la base de datos a `data/empresas.xlsx` "
            "(o indica otra ruta con EMPRESAS_XLSX)."
        )
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        if hoja not in libro.sheetnames:
            raise ValueError(f"El Excel no tiene la hoja «{hoja}» (hojas: {', '.join(libro.sheetnames)}).")
        filas = list(libro[hoja].iter_rows(values_only=True))
    finally:
        libro.close()

    # --- localizar la fila de cabecera (la primera con EMPRESA y SECTOR) ---
    fila_cab = next((i for i, f in enumerate(filas[:FILAS_BUSQUEDA_CABECERA])
                     if {"empresa", "sector"} <= {_norm(c) for c in f if c is not None}), None)
    if fila_cab is None:
        raise ValueError("No se encontró la fila de cabecera (con las columnas EMPRESA y SECTOR) en la hoja «%s»." % hoja)
    cabecera = [_norm(c) for c in filas[fila_cab]]
    columnas, faltan = {}, []
    for clave, (prefijo, requerida) in CABECERAS.items():
        # la 1ª columna "EMPRESA" es el nombre; "SECTOR" exacto (no «DESCRIPCIÓN DEL SECTOR»)
        exacta = clave in ("nombre_empresa", "sector")
        pos = next((i for i, h in enumerate(cabecera) if (h == prefijo if exacta else h.startswith(prefijo))), None)
        if pos is None:
            if requerida:
                faltan.append(prefijo.upper())
        else:
            columnas[clave] = pos
    if faltan:
        raise ValueError("Faltan columnas imprescindibles en el Excel de empresas: " + "; ".join(faltan))

    empresas, usados = [], set()
    for n, fila in enumerate(filas[fila_cab + 1:], start=fila_cab + 2):
        valores = {k: limpiar(fila[i]) if i < len(fila) else None for k, i in columnas.items()}
        nombre = valores.get("nombre_empresa")
        if not nombre:
            continue
        nombre = " ".join(nombre.split())
        identificador = _slug(nombre)
        if identificador in usados:
            identificador = f"{identificador}-{n}"
        usados.add(identificador)
        empresa = {k: valores.get(k) for k in CABECERAS}
        empresa.update(nombre_empresa=nombre, numero_interno=identificador, fila_excel=n)
        # Alias con los nombres de campo que ya usan los informes y la ficha (sin duplicar datos nuevos)
        empresa["paises_interes"] = empresa["zona_geografica_interes"]
        empresa["notas_libres"] = empresa["importes"]
        empresa["observaciones_importantes"] = None
        empresas.append(empresa)
    if not empresas:
        raise ValueError("El Excel de empresas no contiene ninguna empresa.")
    return empresas
