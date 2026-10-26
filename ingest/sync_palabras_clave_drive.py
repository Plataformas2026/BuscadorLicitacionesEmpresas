# -*- coding: utf-8 -*-
"""
sync_palabras_clave_drive.py
----------------------------
Sincroniza el Excel Palabras_clave.xlsx de Google Drive contra la tabla `empresas_palabras_clave` de Supabase. Es lo
ÚNICO que lee la pestaña «Coincidencia Inteligente» (ranking por embeddings + reordenación con GPT) para conocer a las
empresas. Está deliberadamente separado de:
  - BBDD_Empresas.xlsx         -> sync_empresas_drive.py (Directorio y Estadísticas de adjudicación),
  - Seguimiento_Empresas.xlsx  -> sync_licitaciones_informadas_drive.py (Estadísticas de interés).
Nada de lo que se carga aquí toca las tablas `empresas`, `empresas_referencias`, `empresas_hojas_extra` ni
`licitaciones_informadas`, ni al revés.

Se ejecuta una vez al día (.github/workflows/sincronizar_palabras_clave.yml) y solo hace trabajo real si el Excel ha cambiado
en Drive (se compara `modifiedTime`, guardado en `sync_estado`) o si cambió la lógica de este script (VERSION_INGESTA).
Si CUALQUIER carga falla, termina con error y no guarda el estado: la siguiente ejecución lo reintenta.

QUÉ LEE
-------
Hoja «Hoja1» (las hojas ocultas «Hoja2» y «Hoja3» se ignoran), una empresa por fila. La fila de cabecera se busca en las
primeras filas (la que tiene EMPRESA y SECTOR) y las columnas se localizan por el TEXTO de su cabecera, no por su letra:
añadir o mover columnas no rompe nada; si falta una columna imprescindible (EMPRESA, SECTOR, palabras clave extraídas de la
descripción, PRINCIPALES PROYECTOS TIPO) se detiene SIN tocar Supabase. Hay dos columnas llamadas «EMPRESA»: la primera es el
nombre; la segunda (un texto largo) solo se conserva en `datos_excel`.

CABECERA DEL EXCEL                                       -> COLUMNA DE `empresas_palabras_clave`
  EMPRESA (1ª)                                           -> nombre_empresa          [campo 1 del ranking por embeddings]
  SECTOR                                                 -> sector                  [campo 2]
  palabras clave extraidas de LA DESCRIPCIÓN...          -> palabras_clave          [campo 3]
  PRINCIPALES PROYECTOS TIPO                             -> proyectos_tipo          [campo 4]
  DESCRIPCIÓN DEL SECTOR / ESPECIALIDAD EXTRAIDA DE LAS REFERENCIAS / DESCRIPCIÓN BREVE DE LA ACTIVIDAD ->
                                                            descripcion_sector / especialidad_referencias / descripcion_actividad
  PALABRAS CLAVE DE LAS REFERENCIAS / DE LA BBDD / palabras clave unificadas ->
                                                            palabras_clave_referencias / palabras_clave_bbdd / palabras_clave_unificadas
  EXPERIENCIA PAÍSES / ZONA GEOGRÁFICA y PAÍSES DE INTERÉS / ZONAS Y PAISES DESCARTADOS / ÁMBITO GEOGRÁFICO DE OPERACIÓN ->
                                                            experiencia_paises / zona_geografica_interes / zonas_descartadas / ambito_geografico
  TAMAÑO / CERTIFICACIONES / FACTURACIÓN ANUAL / CLIENTES PRINCIPALES / importe mínimo-máximo ->
                                                            tamano / certificaciones / facturacion_anual / clientes_principales / importes
  (fila completa, cabecera real -> valor)                -> datos_excel (jsonb)
El Excel no trae identificador: `clave_empresa` se genera a partir del nombre (slug estable; si dos nombres dieran el mismo
se añade el número de fila). Los valores vacíos o «-» se guardan como NULL: no se inventa nada. Los textos se guardan tal cual
(sin trocear) para no perder frases largas; es la app quien los trata.

CÓMO PROBARLO SIN SUPABASE NI GOOGLE
------------------------------------
    python sync_palabras_clave_drive.py --archivo Palabras_clave.xlsx
Lee el Excel local, enseña un resumen y no escribe nada.

Variables de entorno (modo normal):
    SUPABASE_URL, SUPABASE_SERVICE_KEY
    GOOGLE_SERVICE_ACCOUNT_JSON            -- el mismo secreto que usan los otros dos scripts
    GOOGLE_DRIVE_FILE_ID_PALABRAS_CLAVE    -- File ID de Palabras_clave.xlsx (o FILE_ID_PALABRAS_CLAVE en archivos_drive.py)
Opcional: FORZAR_SINCRONIZACION=true
"""
import io
import json
import os
import re
import sys
import unicodedata

import openpyxl

from archivos_drive import FILE_ID_PALABRAS_CLAVE, file_id_de

CLAVE_SYNC_ESTADO = "palabras_clave_drive_modified_time"
CLAVE_VERSION_INGESTA = "palabras_clave_ingesta_version"
# Súbela cuando cambie la LÓGICA de este script: fuerza una resincronización aunque el Excel no cambie.
VERSION_INGESTA = "2026-10-26.v1"

TABLA = "empresas_palabras_clave"
HOJA = "Hoja1"
FILAS_BUSCANDO_CABECERA = 20
TAMANO_LOTE = 25

# columna de la tabla -> (texto normalizado con el que empieza la cabecera del Excel, imprescindible)
CABECERAS = {
    "nombre_empresa": ("empresa", True),                       # exacta; la 1ª columna con ese nombre
    "sector": ("sector", True),                                # exacta (no «DESCRIPCIÓN DEL SECTOR»)
    "descripcion_sector": ("descripcion del sector", False),
    "especialidad_referencias": ("especialidad extraida de las referencias", False),
    "descripcion_actividad": ("descripcion breve de la actividad", False),
    "palabras_clave_referencias": ("palabras clave de las referencias", False),
    "palabras_clave_bbdd": ("palabras clave de la bbdd", False),
    "palabras_clave_unificadas": ("palabras clave unificadas", False),
    "palabras_clave": ("palabras clave extraidas de la descripcion", True),
    "proyectos_tipo": ("principales proyectos tipo", True),
    "experiencia_paises": ("experiencia paises", False),
    "zona_geografica_interes": ("zona geografica y paises de interes", False),
    "zonas_descartadas": ("zonas y paises descartados", False),
    "tamano": ("tamano", False),
    "certificaciones": ("certificaciones", False),
    "facturacion_anual": ("facturacion anual", False),
    "clientes_principales": ("clientes principales", False),
    "ambito_geografico": ("ambito geografico de operacion", False),
    "importes": ("importe minimo", False),
}
EXACTAS = ("nombre_empresa", "sector")
VACIOS = {"", "-", "--", "—", "n/a", "na", "ninguno", "ninguna", "no consta", "s/d"}


def _norm(texto) -> str:
    texto = unicodedata.normalize("NFD", str(texto or "").casefold())
    return " ".join("".join(c for c in texto if unicodedata.category(c) != "Mn").split())


def _limpiar(valor):
    """Texto limpio (conserva los saltos de línea del Excel) o None si está vacío."""
    if valor is None:
        return None
    lineas = [" ".join(l.split()) for l in str(valor).replace("\r", "\n").split("\n")]
    texto = "\n".join(l for l in lineas if l).strip()
    return None if _norm(texto) in VACIOS else texto


def _slug(nombre: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", _norm(nombre)).strip("-")[:60] or "empresa"


def _coincide(columna: str, cabecera_norm: str) -> bool:
    prefijo = CABECERAS[columna][0]
    return cabecera_norm == prefijo if columna in EXACTAS else cabecera_norm.startswith(prefijo)


# ------------------------------------------------------------------
# Lectura del Excel
# ------------------------------------------------------------------
def leer_registros(origen) -> list:
    """Filas de `empresas_palabras_clave` a partir del Excel (ruta o BytesIO). Lanza RuntimeError si la estructura no es válida."""
    libro = openpyxl.load_workbook(origen, read_only=True, data_only=True)
    try:
        if HOJA not in libro.sheetnames:
            raise RuntimeError(f"El Excel no tiene la hoja «{HOJA}» (hojas: {', '.join(libro.sheetnames)}). No se ha tocado Supabase.")
        filas = list(libro[HOJA].iter_rows(values_only=True))
    finally:
        libro.close()

    fila_cab = next((i for i, f in enumerate(filas[:FILAS_BUSCANDO_CABECERA])
                     if {"empresa", "sector"} <= {_norm(c) for c in f if c is not None}), None)
    if fila_cab is None:
        raise RuntimeError(f"No se encontró la fila de cabecera (con EMPRESA y SECTOR) en «{HOJA}». No se ha tocado Supabase.")

    cabeceras, vistas = {}, {}                     # posición -> cabecera tal cual (las repetidas llevan « (2)»)
    for i, c in enumerate(filas[fila_cab]):
        texto = " ".join(str(c).split()) if c is not None else ""
        if texto:
            vistas[texto] = vistas.get(texto, 0) + 1
            cabeceras[i] = texto if vistas[texto] == 1 else f"{texto} ({vistas[texto]})"

    posicion = {}                                   # columna de la tabla -> posición en la fila (la 1ª que coincida)
    for columna in CABECERAS:
        posicion[columna] = next((i for i, h in cabeceras.items() if _coincide(columna, _norm(h))), None)
    faltan = [CABECERAS[c][0].upper() for c, (_, req) in CABECERAS.items() if req and posicion[c] is None]
    if faltan:
        raise RuntimeError("Faltan columnas imprescindibles en «%s»: %s. No se ha tocado Supabase." % (HOJA, "; ".join(faltan)))

    registros, usados = [], set()
    for n, fila in enumerate(filas[fila_cab + 1:], start=fila_cab + 2):
        def valor(i):
            return _limpiar(fila[i]) if i is not None and i < len(fila) else None

        nombre = valor(posicion["nombre_empresa"])
        if not nombre:
            continue
        nombre = " ".join(nombre.split())
        clave = _slug(nombre)
        if clave in usados:
            clave = f"{clave}-{n}"
        usados.add(clave)
        registro = {c: valor(p) for c, p in posicion.items()}
        registro.update(
            nombre_empresa=nombre, clave_empresa=clave, fila_excel=n,
            datos_excel={h: _limpiar(fila[i]) if i < len(fila) else None for i, h in cabeceras.items()},
        )
        registros.append(registro)
    return registros


# ------------------------------------------------------------------
# Google Drive
# ------------------------------------------------------------------
def obtener_servicio_drive():
    from google.oauth2 import service_account
    from googleapiclient.discovery import build

    credenciales_json = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
    if not credenciales_json:
        raise RuntimeError("Falta la variable de entorno GOOGLE_SERVICE_ACCOUNT_JSON.")
    credenciales = service_account.Credentials.from_service_account_info(
        json.loads(credenciales_json), scopes=["https://www.googleapis.com/auth/drive.readonly"]
    )
    return build("drive", "v3", credentials=credenciales)


def descargar_excel(servicio, file_id: str, mime_type: str) -> io.BytesIO:
    """Sirve tanto para un Google Sheets (se exporta a .xlsx) como para un .xlsx subido a Drive."""
    if mime_type == "application/vnd.google-apps.spreadsheet":
        contenido = servicio.files().export_media(
            fileId=file_id, mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ).execute()
    else:
        contenido = servicio.files().get_media(fileId=file_id).execute()
    return io.BytesIO(contenido)


# ------------------------------------------------------------------
# Supabase
# ------------------------------------------------------------------
def _leer_estado(supabase, clave: str):
    respuesta = supabase.table("sync_estado").select("valor").eq("clave", clave).limit(1).execute()
    return respuesta.data[0]["valor"] if respuesta.data else None


def _guardar_estado(supabase, clave: str, valor: str):
    supabase.table("sync_estado").upsert({"clave": clave, "valor": valor}, on_conflict="clave").execute()


def comprobar_tabla(supabase):
    """Se detiene ANTES de tocar nada si falta la tabla, diciendo qué ejecutar."""
    try:
        supabase.table(TABLA).select("clave_empresa").limit(1).execute()
    except Exception as error:  # noqa: BLE001
        raise RuntimeError(
            f"No existe la tabla `{TABLA}` en Supabase. Ejecuta sql/crear_empresas_palabras_clave.sql en el SQL Editor y vuelve a "
            f"lanzar la sincronización. No se ha modificado nada. Detalle: {error}"
        ) from None


def _escribir_resumen(lineas: list):
    ruta = os.environ.get("GITHUB_STEP_SUMMARY")
    if ruta:
        with open(ruta, "a", encoding="utf-8") as f:
            f.write("### Sincronización de palabras clave\n\n" + "\n".join(f"- {l}" for l in lineas) + "\n")


def _resumen_de(registros: list) -> list:
    return [f"Empresas leídas: {len(registros)}"] + [
        f"Sin {c}: {sum(1 for r in registros if not r.get(c))}" for c in ("palabras_clave", "proyectos_tipo", "zonas_descartadas")
    ]


def ejecutar_sincronizacion():
    print("=" * 100, flush=True)
    print("SINCRONIZACIÓN DE PALABRAS CLAVE — Google Drive -> Supabase", flush=True)
    print("=" * 100, flush=True)

    file_id = file_id_de("GOOGLE_DRIVE_FILE_ID_PALABRAS_CLAVE", FILE_ID_PALABRAS_CLAVE)
    if not file_id:
        raise RuntimeError(
            "Falta el File ID de Palabras_clave.xlsx: define el secreto GOOGLE_DRIVE_FILE_ID_PALABRAS_CLAVE en GitHub "
            "o rellena FILE_ID_PALABRAS_CLAVE en ingest/archivos_drive.py (pasos en README.md, sección «Palabras clave»)."
        )

    from common import obtener_cliente_supabase, subir_en_lotes

    supabase = obtener_cliente_supabase()
    servicio = obtener_servicio_drive()
    comprobar_tabla(supabase)

    metadata = servicio.files().get(fileId=file_id, fields="modifiedTime, name, mimeType").execute()
    modificado_en = metadata["modifiedTime"]
    print(f"Fichero: {metadata.get('name')}  ·  Última modificación en Drive: {modificado_en}", flush=True)

    ultimo_conocido = _leer_estado(supabase, CLAVE_SYNC_ESTADO)
    version_guardada = _leer_estado(supabase, CLAVE_VERSION_INGESTA)
    motivo = None
    if os.environ.get("FORZAR_SINCRONIZACION", "").strip().lower() in {"1", "true", "yes", "si", "sí"}:
        motivo = "se ha pedido forzar la sincronización"
    elif version_guardada != VERSION_INGESTA:
        motivo = f"ha cambiado la lógica de ingesta ({version_guardada or 'sin versión previa'} -> {VERSION_INGESTA})"
    if ultimo_conocido == modificado_en and not motivo:
        print("No hay cambios desde la última sincronización. Nada que hacer.", flush=True)
        return

    print(f"Descargando ({motivo or 'cambio detectado o primera sincronización'})...", flush=True)
    buffer_excel = descargar_excel(servicio, file_id, metadata.get("mimeType", ""))

    # Se lee TODO antes de tocar Supabase: si el Excel no se puede leer, la base queda intacta.
    registros = leer_registros(buffer_excel)
    if not registros:
        raise RuntimeError("No se ha leído ninguna empresa del Excel. Se aborta sin tocar Supabase.")
    resumen = _resumen_de(registros)
    print("\n".join(resumen), flush=True)

    errores = []
    subidas = subir_en_lotes(supabase, TABLA, "clave_empresa", registros, tamano_lote=TAMANO_LOTE)
    resumen.append(f"Empresas sincronizadas: {subidas}/{len(registros)}")
    if subidas != len(registros):
        errores.append(f"solo se han subido {subidas} de {len(registros)} empresas")

    actuales = {r["clave_empresa"] for r in registros}
    existentes = supabase.table(TABLA).select("id, clave_empresa").execute().data or []
    a_borrar = [f["id"] for f in existentes if f["clave_empresa"] not in actuales]
    for i in range(0, len(a_borrar), 100):
        supabase.table(TABLA).delete().in_("id", a_borrar[i:i + 100]).execute()
    if a_borrar:
        resumen.append(f"Empresas retiradas (ya no están en el Excel): {len(a_borrar)}")
        print(resumen[-1], flush=True)

    if errores:
        _escribir_resumen(resumen + [f"❌ {e}" for e in errores])
        raise RuntimeError("Sincronización INCOMPLETA (se reintentará en la próxima ejecución): " + " | ".join(errores))

    _guardar_estado(supabase, CLAVE_SYNC_ESTADO, modificado_en)
    _guardar_estado(supabase, CLAVE_VERSION_INGESTA, VERSION_INGESTA)
    print("Estado de sincronización actualizado.", flush=True)
    _escribir_resumen(resumen + ["✅ Sincronización completa"])


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "--archivo":
        datos = leer_registros(sys.argv[2])
        print("\n".join(_resumen_de(datos)))
        for muestra in datos[:2]:
            print({k: (str(v)[:60] if v is not None else None) for k, v in muestra.items() if k != "datos_excel"})
    else:
        ejecutar_sincronizacion()
